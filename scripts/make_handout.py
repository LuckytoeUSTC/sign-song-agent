"""One handout entry point: production JSON or a selected DOCX copy."""
from pathlib import Path
import argparse
import copy
import json
import math
from zipfile import ZipFile
from lxml import etree

from docx import Document
from docx.image.image import Image
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

TOOL_VERSION = "3.1.1"
DEFAULT_LAYOUT = dict(font_name='SimSun', font_size=11, title_size=16,
                      heading_size=13, image_height=.79, margin_horizontal=2,
                      margin_vertical=1.8, line_spacing=1.15)


def text_runs(p, text, font_name='SimSun', font_size=11):
    """Render the transcription markers and run boundaries."""
    italic = strike = False
    index = 0
    while index < len(text):
        if text.startswith('~~', index):
            strike = not strike
            index += 2
            continue
        if text[index] == '*':
            italic = not italic
            index += 1
            continue
        if text.startswith('++', index):
            token = '++'
        elif text[index] == '~':
            token = '~'
        else:
            end = index + 1
            while end < len(text) and text[end] not in '*~' and not text.startswith('++', end):
                end += 1
            token = text[index:end]
        run = p.add_run(token)
        run.italic = italic and token != '++'
        run.font.strike = strike
        if token == '++':
            ics = OxmlElement('w:iCs')
            ics.set(qn('w:val'), '0')
            run._element.get_or_add_rPr().append(ics)
        run.font.name = 'Arial' if token == '~' else font_name
        run.font.size = Pt(font_size)
        run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), font_name)
        index += len(token)
    if italic or strike:
        raise ValueError('未闭合的斜体或删除线标记：' + text)


STYLE_FIELDS = {'font_name', 'font_size', 'space_before', 'space_after', 'line_spacing'}


def validate_paragraph_style(options):
    if not isinstance(options, dict) or set(options) - STYLE_FIELDS:
        raise ValueError('段落格式只支持字体、字号、段前段后及行距')
    for key, value in options.items():
        if key == 'font_name':
            if not isinstance(value, str) or not value.strip():
                raise ValueError('字体须为非空文字')
        elif type(value) not in (int, float) or not math.isfinite(value) or (
                value < 0 if key.startswith('space_') else value <= 0):
            raise ValueError('无效段落格式：' + key)


def apply_paragraph_style(p, options):
    for key in ('space_before', 'space_after', 'line_spacing'):
        if key in options:
            setattr(p.paragraph_format, key,
                    options[key] if key == 'line_spacing' else Pt(options[key]))
    for run in p.runs:
        if 'font_name' in options and run.text != '~':
            run.font.name = options['font_name']
            run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), options['font_name'])
        if 'font_size' in options:
            run.font.size = Pt(options['font_size'])


def pad_reference_row(p, padding):
    """Reserve vertical room only for a selected row containing dotted pictures."""
    from docx.enum.text import WD_LINE_SPACING
    heights = p._p.xpath('.//wp:inline/wp:extent/@cy')
    lines = p._p.xpath('.//pic:spPr/a:ln[a:prstDash]')
    if not heights or not lines:
        return
    required = max(int(h) for h in heights) / 12700 + padding
    current = p.paragraph_format.line_spacing
    if isinstance(current, int):
        required = max(required, current / 12700)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    p.paragraph_format.line_spacing = Pt(required)
    for line in lines:
        line.set('algn', 'in')


def copy_reference(source, output, supplied):
    """Patch only document.xml; all media, styles, headers and other parts stay intact."""
    document = Document(source)
    forbidden = set(supplied) - {'paragraph_overrides', 'reference_frame_rows',
                                'reference_frame_padding_pt', 'content_styles'}
    if forbidden:
        raise ValueError('docx副本只接受逐段覆盖和指定参考图行；全局参数用于新稿')
    if supplied.get('content_styles'):
        raise ValueError('docx副本请用paragraph_overrides明确选择段落，不推测角色')
    overrides = supplied.get('paragraph_overrides', {})
    rows = supplied.get('reference_frame_rows', [])
    for number in list(overrides) + rows:
        if int(number) > len(document.paragraphs):
            raise ValueError('段落序号超出原件：' + str(number))
    for number, options in overrides.items():
        apply_paragraph_style(document.paragraphs[int(number)-1], options)
    for number in rows:
        pad_reference_row(document.paragraphs[number-1], supplied.get('reference_frame_padding_pt', 6))
    with ZipFile(source) as original, ZipFile(output, 'w') as result:
        for info in original.infolist():
            data = original.read(info.filename)
            if info.filename == 'word/document.xml' and (overrides or rows):
                data = etree.tostring(document._element, xml_declaration=True,
                                      encoding='UTF-8', standalone=True)
            result.writestr(info, data)


def read_layout(path=None, height=None):
    settings = DEFAULT_LAYOUT.copy()
    settings['reference_frame_padding'] = False
    settings.update(content_styles={}, paragraph_overrides={}, reference_frame_rows=[],
                    reference_frame_padding_pt=6)
    if path:
        supplied = json.loads(Path(path).read_text(encoding='utf-8'))
        unknown = set(supplied) - set(settings) - {'image_space_before', 'image_space_after', 'image_line_spacing'}
        if unknown:
            raise ValueError('未知版式参数：' + ', '.join(sorted(unknown)))
        settings.update(supplied)
    for key, value in settings.items():
        if key in ('content_styles', 'paragraph_overrides'):
            if not isinstance(value, dict):
                raise ValueError(key + '须为对象')
            for name, options in value.items():
                if key == 'paragraph_overrides' and (not str(name).isdigit() or int(name) < 1):
                    raise ValueError('段落序号从1开始')
                validate_paragraph_style(options)
            continue
        if key == 'reference_frame_rows':
            if not isinstance(value, list) or any(type(n) is not int or n < 1 for n in value):
                raise ValueError('reference_frame_rows须为从1开始的段落序号数组')
            continue
        if key in ('image_space_before', 'image_space_after', 'reference_frame_padding_pt'):
            if type(value) not in (int, float) or not math.isfinite(value) or value < 0:
                raise ValueError(key + '须为非负有限数')
            continue
        if key == 'reference_frame_padding':
            if not isinstance(value, bool):
                raise ValueError('reference_frame_padding须为true或false')
            continue
        if key != 'font_name' and (not isinstance(value, (int, float)) or value <= 0):
            raise ValueError('版式参数必须是正数：' + key)
    if height is not None:
        settings['image_height'] = height
    if not .1 <= settings['image_height'] <= 3:
        raise ValueError('图片高度应为0.1至3英寸')
    return settings


def read_content(source, input_format='json'):
    """Read the current production JSON schema."""
    source = Path(source).resolve()
    text = source.read_text(encoding='utf-8')
    if input_format != 'json':
        raise ValueError('完整版制作输入仅支持JSON；既有Word请使用docx副本模式')
    data = json.loads(text)
    if not isinstance(data, dict) or data.get('version') != 1 or not isinstance(data.get('blocks'), list):
        raise ValueError('制作JSON须含version: 1和blocks数组')
    # Resolve only previous IDs: this makes repeat cycles impossible.
    resolved = []
    by_id = {}
    for original in data['blocks']:
        if not isinstance(original, dict):
            raise ValueError('每个制作块须为对象')
        block = copy.deepcopy(original)
        if 'repeat_of' in block:
            target = block['repeat_of']
            if target not in by_id:
                raise ValueError('重复段须引用前面已定义的id：' + str(target))
            inherited = copy.deepcopy(by_id[target])
            for key in ('lyrics', 'transcription', 'images', 'heading', 'text', 'level'):
                if key in block:
                    raise ValueError('重复段不另存一份正文或图片；修改原块或取消repeat_of：' + key)
            inherited.pop('id', None)
            inherited.update(block)
            block = inherited
            if block.get('omit_images', False):
                block['images'] = []
        elif 'omit_images' in block:
            raise ValueError('omit_images仅用于repeat_of重复段')
        if 'id' in block:
            if not isinstance(block['id'], str) or not block['id'] or block['id'] in by_id:
                raise ValueError('制作块id须为唯一的非空字符串')
            by_id[block['id']] = copy.deepcopy(block)
        resolved.append(block)
    return dict(input_format='json', base=source.parent, blocks=resolved)


def check_image(item, base):
    """Check file/rectangle facts, not whether a hand action is natural."""
    if not isinstance(item, dict) or not isinstance(item.get('path'), str):
        raise ValueError('每张选图须有path；不能用词条名代替已选文件')
    path = (base / item['path']).resolve()
    if not path.is_file():
        raise FileNotFoundError(f'图片不存在：{path}')
    image = Image.from_file(str(path))
    crop = item.get('crop')
    if crop is not None:
        if not isinstance(crop, list) or len(crop) != 4 or any(
                not isinstance(v, (int, float)) or not math.isfinite(v) for v in crop):
            raise ValueError('crop须为原图像素坐标[left, top, right, bottom]')
        left, top, right, bottom = crop
        if not (0 <= left < right <= image.px_width and 0 <= top < bottom <= image.px_height):
            raise ValueError('裁图范围越界或为空：' + str(path))
    return path, image


def check_content(content):
    for block in content['blocks']:
        for key in ('lyrics', 'transcription', 'text', 'heading', 'note', 'record', 'role'):
            if key in block and not isinstance(block[key], str):
                raise ValueError(key + '须为文字')
        if 'heading' in block and any(k in block for k in ('lyrics', 'transcription', 'text', 'images', 'note')):
            raise ValueError('标题块不同时包含正文、图片或短注释')
        if 'text' in block and any(k in block for k in ('lyrics', 'transcription')):
            raise ValueError('正文text与歌词/转写二选一，避免并行数据')
        if 'heading' in block and block.get('level', 1) not in (1, 2, 3):
            raise ValueError('标题level须为1、2或3')
        if not any(key in block for key in ('heading', 'text', 'lyrics', 'transcription')):
            raise ValueError('制作块须有标题、正文或歌词转写')
        if not isinstance(block.get('images', []), list):
            raise ValueError('images须为数组')
        if 'omit_images' in block and not isinstance(block['omit_images'], bool):
            raise ValueError('omit_images须为true或false')
        for item in block.get('images', []):
            check_image(item, content['base'])
            for key in ('entry', 'sense', 'source', 'action', 'note', 'record', 'alt'):
                if key in item and not isinstance(item[key], str):
                    raise ValueError('图片' + key + '须为文字')
            if 'reference' in item and not isinstance(item['reference'], bool):
                raise ValueError('reference须为true或false')
    return content


def new_document(settings):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(settings['margin_horizontal'])
    section.top_margin = section.bottom_margin = Cm(settings['margin_vertical'])
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
        style = document.styles[name]
        style.font.name = settings['font_name']
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), settings['font_name'])
    for style in document.styles:
        for border in style.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    document.styles['Normal'].font.size = Pt(settings['font_size'])
    return document


def add_heading(document, block, settings):
    level = block.get('level', 1)
    p = document.add_paragraph(block['heading'], style='Title' if level == 1 else 'Heading ' + str(level - 1))
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    for run in p.runs:
        run.font.size = Pt(settings['title_size'] if level == 1 else settings['heading_size'])
    return p


def add_paragraph(document, settings, images=False, keep_with_next=False):
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(7 if images else 5)
    p.paragraph_format.line_spacing = 1.25 if images else settings['line_spacing']
    p.paragraph_format.keep_together = True
    if keep_with_next:
        p.paragraph_format.keep_with_next = True
    if images:
        p.paragraph_format.space_before = Pt(4)
        if 'image_space_before' in settings:
            p.paragraph_format.space_before = Pt(settings['image_space_before'])
        if 'image_space_after' in settings:
            p.paragraph_format.space_after = Pt(settings['image_space_after'])
        if 'image_line_spacing' in settings:
            p.paragraph_format.line_spacing = settings['image_line_spacing']
    return p


def add_picture(p, item, base, settings):
    path, image = check_image(item, base)
    shape = p.add_run().add_picture(str(path), height=Inches(settings['image_height']))
    # Alt/entry metadata is never turned into a visible caption.
    shape._inline.docPr.set('descr', item.get('alt', item.get('entry', '')))
    if 'crop' in item:
        left, top, right, bottom = item['crop']
        shape.width = int(shape.height * (right - left) / (bottom - top))
        rect = OxmlElement('a:srcRect')
        for key, fraction in dict(l=left / image.px_width, t=top / image.px_height,
                                  r=1 - right / image.px_width, b=1 - bottom / image.px_height).items():
            rect.set(key, str(round(fraction * 100000)))
        fill = shape._inline.xpath('.//pic:blipFill')[0]
        stretch = fill.find(qn('a:stretch'))
        if stretch is None:
            fill.append(rect)
        else:
            fill.insert(list(fill).index(stretch), rect)
    if item.get('reference', False):
        for effect in shape._inline.xpath('./wp:effectExtent'):
            for key, value in dict(t='25400', b='25400', l='12700', r='12700').items():
                effect.set(key, value)
        props = shape._inline.xpath('.//pic:spPr')[0]
        line = OxmlElement('a:ln')
        line.set('w', '12700')
        if settings.get('reference_frame_padding', False):
            line.set('algn', 'in')
        fill = OxmlElement('a:solidFill')
        color = OxmlElement('a:srgbClr')
        color.set('val', '000000')
        fill.append(color)
        line.append(fill)
        dash = OxmlElement('a:prstDash')
        dash.set('val', 'sysDot')
        line.append(dash)
        props.append(line)
    return shape


def render_content(content, settings):
    document = new_document(settings)
    write = lambda p, value: text_runs(p, value, settings['font_name'], settings['font_size'])
    for block in content['blocks']:
        if 'heading' in block:
            add_heading(document, block, settings)
            continue
        images = block.get('images', [])
        text = block.get('text', block.get('lyrics', ''))
        if 'transcription' in block:
            text += '（' + block['transcription'] + '）'
        p = add_paragraph(document, settings, keep_with_next=bool(images))
        roles = settings.get('content_styles', {})
        if 'transcription' in block and any(k in roles for k in ('lyrics', 'transcription')):
            options = roles.get('lyrics', {})
            text_runs(p, block.get('lyrics', ''), options.get('font_name', settings['font_name']),
                      options.get('font_size', settings['font_size']))
            options = roles.get('transcription', {})
            text_runs(p, '（' + block['transcription'] + '）',
                      options.get('font_name', settings['font_name']),
                      options.get('font_size', settings['font_size']))
        else:
            write(p, text)
        options = roles.get(block.get('role', 'body'), {})
        if 'transcription' in block and any(k in roles for k in ('lyrics', 'transcription')):
            options = {k: v for k, v in options.items() if k not in ('font_name', 'font_size')}
        apply_paragraph_style(p, options)
        if images:
            picture_row = add_paragraph(document, settings, images=True)
            for item in images:
                add_picture(picture_row, item, content['base'], settings)
                if item.get('note'):
                    start = len(picture_row.runs)
                    write(picture_row, '〔' + item['note'] + '〕')
                    options = settings.get('content_styles', {}).get('note', {})
                    for run in picture_row.runs[start:]:
                        if 'font_name' in options:
                            run.font.name = options['font_name']
                            run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), options['font_name'])
                        if 'font_size' in options:
                            run.font.size = Pt(options['font_size'])
                write(picture_row, ' ')
            if settings.get('reference_frame_padding'):
                pad_reference_row(picture_row, settings['reference_frame_padding_pt'])
        if block.get('note'):
            note_row = add_paragraph(document, settings)
            write(note_row, block['note'])
            apply_paragraph_style(note_row, settings.get('content_styles', {}).get('note', {}))
    return document


def make_handout(source, output, height=None, layout=None, input_format='json'):
    output = Path(output)
    if output.exists():
        raise ValueError('输出文件已存在；请选择新文件名')
    settings = read_layout(layout, height)
    if input_format == 'docx':
        if height is not None:
            raise ValueError('docx副本保留原图尺寸，不接受--height')
        supplied = json.loads(Path(layout).read_text(encoding='utf-8')) if layout else {}
        output.parent.mkdir(parents=True, exist_ok=True)
        copy_reference(source, output, supplied)
        return output
    content = check_content(read_content(source, input_format))
    document = render_content(content, settings)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return output


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--version', action='version', version=TOOL_VERSION)
    ap.add_argument('source', type=Path, nargs='?')
    ap.add_argument('--out', type=Path)
    ap.add_argument('--height', type=float)
    ap.add_argument('--layout', type=Path, help='JSON layout settings')
    ap.add_argument('--input-format', choices=('json', 'docx'), default='json')
    args = ap.parse_args()
    if not args.source or not args.out:
        ap.error('制作需要source与--out')
    make_handout(args.source, args.out, args.height, args.layout, args.input_format)
    print('生成', args.out, '；请逐页预览后使用。')


if __name__ == '__main__':
    main()
