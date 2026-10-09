"""Read-only Word checks: confirmed errors, warnings and human review items."""
from pathlib import Path
from collections import Counter
import argparse
import io
import re
from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.image.image import Image
from docx.oxml.ns import qn


def effective(r, p, attribute):
    direct = getattr(r.font, attribute)
    if direct is not None:
        return direct
    for style in [r.style, p.style]:
        while style:
            value = getattr(style.font, attribute)
            if value is not None:
                return value
            style = style.base_style
    return None


def _crop_values(drawing):
    found = drawing.xpath('.//a:srcRect')
    return tuple(int(found[0].get(k, '0')) for k in ('l', 't', 'r', 'b')) if found else (0, 0, 0, 0)


def _selected_bytes(drawing, document):
    blips = drawing.xpath('.//a:blip')
    part = document.part.related_parts.get(blips[0].get(qn('r:embed'))) if blips else None
    return part.blob if part is not None else None


def inspect_docx(path, height=None, source=None, input_format='json'):
    """No writes. Missing optional production fields are not errors."""
    d = Document(path)
    errors, warnings = [], []
    manual = ['逐页查看导出PDF或页面图，核对分页、图框、文字和留白。',
              '手形、箭头是否被裁掉、借图是否适当及动作自然度须看原图并由老师试打；静态检查不能确认。']
    fonts, sizes = Counter(), Counter()
    image_headers = {}  # Per-document cache; never retained across file changes.
    transcriptions, images = [], 0
    width = min(s.page_width - s.left_margin - s.right_margin for s in d.sections)
    paragraphs = list(d.paragraphs) + [p for t in d.tables for row in t.rows for c in row.cells for p in c.paragraphs]
    for i, p in enumerate(paragraphs, 1):
        text = p.text
        if not text.startswith('【') and '（' in text and '/' in text:
            transcriptions.append(text.split('（', 1)[1].rsplit('）', 1)[0])
        flags = []
        for r in p.runs:
            flags.extend([bool(effective(r, p, 'italic'))] * len(r.text))
            if r.text:
                fonts[effective(r, p, 'name') or '未显式指定'] += len(r.text)
                size = effective(r, p, 'size')
                sizes[round(size.pt, 1) if size else '未显式指定'] += len(r.text)
        for match in re.finditer(r'\+\+', text):
            if any(flags[match.start():match.end()]):
                errors.append(f'段落{i}：++含斜体')
        if '**' in text or '\\n' in text:
            warnings.append(f'段落{i}：可能残留Markdown或转义字符')
        for drawing in p._p.xpath('.//w:drawing'):
            images += 1
            extents = drawing.xpath('.//wp:extent')
            if drawing.xpath('./wp:anchor'):
                warnings.append(f'段落{i}：浮动图片，检查位置')
            if not extents:
                continue
            h, w = int(extents[0].get('cy')), int(extents[0].get('cx'))
            if height and abs(h / 914400 - height) > .002:
                warnings.append(f'段落{i}：图高{h / 914400:.3f}英寸，检查是否为照片等例外')
            if w > width:
                warnings.append(f'段落{i}：单图宽度超过正文宽度')
            if p.paragraph_format.line_spacing_rule == WD_LINE_SPACING.EXACTLY:
                line = p.paragraph_format.line_spacing
                if line and line < h:
                    warnings.append(f'段落{i}：固定行距小于图高，可能截框')
            # Compare stored image geometry to the visible crop; no semantic claims.
            blips = drawing.xpath('.//a:blip')
            if blips and h > 0:
                rel = blips[0].get(qn('r:embed'))
                if rel and rel in d.part.related_parts:
                    try:
                        if rel not in image_headers:
                            image_headers[rel] = Image.from_file(io.BytesIO(d.part.related_parts[rel].blob))
                        image = image_headers[rel]
                        left, top, right, bottom = _crop_values(drawing)
                        visible_w = 1 - (left + right) / 100000
                        visible_h = 1 - (top + bottom) / 100000
                        if visible_w <= 0 or visible_h <= 0:
                            errors.append(f'段落{i}：图片裁图范围为空')
                        else:
                            expected = (image.px_width / image.horz_dpi) / (image.px_height / image.vert_dpi) * visible_w / visible_h
                            # JSON pixel crops use square-pixel display; unusual DPI needs review.
                            pixel_expected = image.px_width / image.px_height * visible_w / visible_h
                            if min(abs(w / h / expected - 1), abs(w / h / pixel_expected - 1)) > .03:
                                warnings.append(f'段落{i}：图片显示比例与原图及裁图不一致，核对是否有意变形')
                    except (ValueError, KeyError, ZeroDivisionError, TypeError):
                        warnings.append(f'段落{i}：此图片类型未完成比例检查，人工核对')
    if source:
        from make_handout import read_content, check_content, render_content, DEFAULT_LAYOUT
        content = check_content(read_content(source, input_format))
        expected = render_content(content, DEFAULT_LAYOUT)
        actual_text = [p.text for p in d.paragraphs]
        if actual_text != [p.text for p in expected.paragraphs]:
            errors.append('正文或必要短注释与所指定制作输入不一致；检查是否检查了正确版本')
        expected_drawings = expected._element.xpath('.//w:drawing')
        if images != len(expected_drawings):
            errors.append('图片数量与所指定输入的重复引用/省图关系不一致')
        actual_drawings = d._element.xpath('.//w:drawing')
        for index, (actual, wanted) in enumerate(zip(actual_drawings, expected_drawings), 1):
            if _crop_values(actual) != _crop_values(wanted):
                errors.append(f'图{index}：裁图范围与制作输入不一致')
            if _selected_bytes(actual, d) != _selected_bytes(wanted, expected):
                errors.append(f'图{index}：实际图片不是指定的选图文件')
            actual_reference = bool(actual.xpath('.//a:prstDash[@val="sysDot"]'))
            wanted_reference = bool(wanted.xpath('.//a:prstDash[@val="sysDot"]'))
            if actual_reference != wanted_reference:
                errors.append(f'图{index}：参考图框与制作输入不一致')
    return dict(path=str(path), errors=errors, warnings=warnings, manual=manual,
                fonts=dict(fonts), sizes=dict(sizes), transcriptions=transcriptions, images=images)


def print_findings(result):
    print(result['path'], '\n转写', len(result['transcriptions']), '处；图片', result['images'],
          '幅\n字体分布', result['fonts'], '\n字号分布', result['sizes'])
    for label, key in (('错误', 'errors'), ('警告', 'warnings'), ('人工核对', 'manual')):
        print(label + '：')
        print('\n'.join(result[key]) if result[key] else '无')


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('files', nargs='+', type=Path)
    ap.add_argument('--height', type=float)
    ap.add_argument('--source', type=Path, help='optional approved production input')
    ap.add_argument('--input-format', choices=('json',), default='json')
    args = ap.parse_args()
    if args.source and len(args.files) != 1:
        ap.error('--source只核对一个明确对应的文件')
    results = [inspect_docx(path, args.height, args.source, args.input_format) for path in args.files]
    for result in results:
        print_findings(result)
    values = [result['transcriptions'] for result in results]
    if len(values) == 2:
        if values[0] == values[1]:
            print('两版正文转写文字一致。')
        else:
            print('警告：两版正文转写有差异：')
            from itertools import zip_longest
            for i, (x, y) in enumerate(zip_longest(*values), 1):
                if x != y:
                    print(i, repr(x), '<>', repr(y))
    return 2 if any(result['errors'] for result in results) else 0


if __name__ == '__main__':
    raise SystemExit(main())
