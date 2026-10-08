"""Synthetic production/update regression tests; never read real song data."""
import base64
import hashlib
import io
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

SCRIPTS = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SCRIPTS))
import make_handout as maker
import handout_update as updater
from audit_docx import inspect_docx, audit
from docx import Document
from docx.shared import Pt

PNG = 'iVBORw0KGgoAAAANSUhEUgAAAPAAAAB4CAIAAABD1OhwAAADbklEQVR4nO3d4W6qQBCGYabp/d/yNNHEUFFhhd2Z+fZ9fp5U3aNvp4CI5u4LoOInegHAlQgaUggaUggaUggaUggaUggaUggaUggaUggaUggaUggaUggaUggaUggaUggaUggaUggaUggaUn6bftrMlvr4GOXXr2/+p64taOixliG1/uGccRP0pOz0H9vHPaQq+8ugU/0fptpe6vo8+N7L+vK2939MkgQTeiL2KsemENc//HRvSbLmKMe8NfvN13f48ubhfwaZ0JOmvFzkflfrh4gd1UzouWr2c1P5ne3dRo1qglZm/4+y9Z6aTw8R0jRBy7KgY8axTRO0Jgt9BySwaYIWZAnez4tqmqDVWIKaA5smaCmWpuaopglahyWrOaRpghbkaWoevx6CFpHz3LeHx6p6D2mCVhB+BkWe1RK0FE85nkeujaDLqzWee6+ZoHV44vE8bIUEXVvyfcHxe4cEDSkEXVi58TxgSBM0pBA0pBB0VUW3N3pvdRA0pBA0pBB0SaW3N7pudRA0pBC07PkScyLoozXTdAkEvWN7kStkRtA7klzhSmyPsN9+IUGXbBrvEPQhNF0FQR9F0yUQdAOazm/GC55fuBFsZtX3zMQwoU+h5mwI+hSOeGQz4ybHmbG6/YaHK1aEyzChG1BzfgR9FDWXQNCHUHMVBL2Pmgsh6Ho1D7s0bcWzrAh6x/q5zlAzPiPoffeOqbkEgj6EmqsgaEgh6JIE9gutz+duCBpSCBpSCLqq0lsd1u1zvgQNKQQNKQRdWNGtDut5XRGChhSCrq3ckLbOl30iaB2WvukBKyTo8iqeZ+Ld1kzQUizxkB6zNoJWUGtIe8/VErSI5HuHNuoSwAQtyJI1PXI9BK3DV8MvT9PrlQzYNCJoKZ6s6cE1E7QgT9P0+JoJWlOGpkNqJmhZHtp0VM0ErcyDmg6smaDnato6Z/30ECFv93CUY7rvhbEOWW/vNurNyxkveD4bv7W1/UrcS5rb/nrEvg/PhJ6Fbzo7Oa1f3jz8rBIm9NSjemnch/vwCxCe8l3bt5JlePPpvCRPfSy77qVM9XwyoSflp8/OS9XxA0HPzlsOV+eMeI0vQoUUjnJACkFDCkFDCkFDCkFDCkFDCkFDCkFDCkFDCkFDCkFDCkFDCkFDCkFDCkFDCkFDCkFDCkFjUfIHhfVyD7qKb7kAAAAASUVORK5CYII='
LEGACY = '# 合成兼容样稿\n\n## 小标题\n测试句（*动作①++*/~~省略②~~/~轻声~）\n![内部词条](synthetic.png "参考") 短注释\n第三行（对象/方向）\n'
LAYOUT = dict(font_name='SimSun', font_size=12, title_size=18, heading_size=14,
              image_height=.6, margin_horizontal=2.3, margin_vertical=2, line_spacing=1.2)


def sample_data():
    return {'version': 1, 'blocks': [
        {'heading': '合成制作样稿'},
        {'id': 'verse', 'lyrics': '测试画面', 'transcription': '*动作①++*/~~省略②~~',
         'record': '长说明不打印', 'images': [
             {'path': 'synthetic.png', 'entry': '内部词条不打印', 'sense': '义项不打印',
              'source': '来源不打印', 'action': '动作部分不打印', 'record': '图长说明不打印',
              'alt': '替代文本不打印', 'note': '方向调整', 'reference': True,
              'crop': [0, 0, 120, 120]},
             {'path': 'synthetic.png', 'entry': '无标签图', 'note': '双主体位置'}]},
        {'repeat_of': 'verse', 'omit_images': True, 'note': '重复段沿用前图'},
        {'repeat_of': 'verse'}]}


class ProductionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='sign-production-tests-')
        self.root = Path(self.temp.name)
        (self.root / 'synthetic.png').write_bytes(base64.b64decode(PNG))
        (self.root / 'legacy.md').write_text(LEGACY, encoding='utf-8')
        (self.root / 'layout.json').write_text(json.dumps(LAYOUT), encoding='utf-8')

    def tearDown(self):
        self.temp.cleanup()

    def run_maker(self, source, output, *args):
        result = subprocess.run([sys.executable, str(SCRIPTS / 'make_handout.py'), str(source),
                                 '--out', str(output), *args], capture_output=True)
        self.assertEqual(result.returncode, 0, result.stderr.decode('utf-8', errors='replace'))

    def test_old_markdown_command_and_default_layout_are_identical(self):
        output = self.root / 'default.docx'
        self.run_maker(self.root / 'legacy.md', output)
        with ZipFile(output) as z:
            self.assertEqual(hashlib.sha256(z.read('word/document.xml')).hexdigest(),
                             'f4c5da8ca607377914039db3e6045753ca370aa6bd3af7b37af76a57d9e151d0')
            self.assertEqual(hashlib.sha256(z.read('word/styles.xml')).hexdigest(),
                             '4b9181172feb2e7b974ef1c7565296ef53595a9796e9c67a36484225f97f3890')
        self.assertFalse(inspect_docx(output)['errors'])
        self.assertIsInstance(audit(output), list)  # Old callable contract.

    def test_old_layout_fields_and_height_override_are_identical(self):
        output = self.root / 'layout.docx'
        self.run_maker(self.root / 'legacy.md', output, '--layout', str(self.root / 'layout.json'), '--height', '.65')
        with ZipFile(output) as z:
            self.assertEqual(hashlib.sha256(z.read('word/document.xml')).hexdigest(),
                             '0e0b03161202cec04f1e4b3e642ba8a359d3fbbe36242aa42f5b3ede32898d28')
            self.assertEqual(hashlib.sha256(z.read('word/styles.xml')).hexdigest(),
                             '9691b550666a687b6c3a7bfdcfe81c12c4b8929359a9e712fa3549e38c67a4b3')

    def test_json_notes_crop_repeat_and_markers(self):
        source = self.root / 'content.json'
        source.write_text(json.dumps(sample_data(), ensure_ascii=False), encoding='utf-8')
        before = (self.root / 'synthetic.png').read_bytes()
        output = self.root / 'json.docx'
        self.run_maker(source, output, '--input-format', 'json')
        d = Document(output)
        visible = '\n'.join(p.text for p in d.paragraphs)
        for hidden in ('内部词条不打印', '义项不打印', '来源不打印', '动作部分不打印',
                       '长说明不打印', '替代文本不打印', '无标签图'):
            self.assertNotIn(hidden, visible)
        self.assertEqual(visible.count('测试画面（动作①++/省略②）'), 3)
        self.assertEqual(visible.count('〔方向调整〕'), 2)
        self.assertEqual(visible.count('〔双主体位置〕'), 2)
        self.assertIn('重复段沿用前图', visible)
        self.assertEqual(len(d.inline_shapes), 4)
        self.assertAlmostEqual(d.inline_shapes[0].width / d.inline_shapes[0].height, 1, places=4)
        self.assertAlmostEqual(d.inline_shapes[1].width / d.inline_shapes[1].height, 2, places=4)
        self.assertEqual(before, (self.root / 'synthetic.png').read_bytes())
        with ZipFile(output) as z:
            embedded = next(n for n in z.namelist() if n.startswith('word/media/'))
            self.assertEqual(z.read(embedded), before)
            self.assertIn(b'sysDot', z.read('word/document.xml'))
        flags = [(r.text, r.italic, r.font.strike) for p in d.paragraphs for r in p.runs]
        self.assertIn(('动作①', True, False), flags)
        self.assertIn(('++', False, False), flags)
        self.assertIn(('省略②', False, True), flags)
        result = inspect_docx(output, source=source)
        self.assertFalse(result['errors'], result)
        self.assertFalse(result['warnings'], result)
        d.paragraphs[2].runs[-2].text = '〔误改短注释〕'
        altered = self.root / 'altered.docx'
        d.save(altered)
        self.assertTrue(inspect_docx(altered, source=source)['errors'])

    def test_reference_frame_padding_is_explicit_and_keeps_original_pixels(self):
        from docx.enum.text import WD_LINE_SPACING
        layout = self.root / 'frame.json'
        layout.write_text('{"reference_frame_padding": true}', encoding='utf-8')
        output = self.root / 'frame.docx'
        maker.make_handout(self.root / 'legacy.md', output, layout=layout)
        d = Document(output)
        image_row = next(p for p in d.paragraphs if p._p.xpath('.//w:drawing'))
        self.assertEqual(image_row.paragraph_format.line_spacing_rule, WD_LINE_SPACING.AT_LEAST)
        self.assertAlmostEqual(image_row.paragraph_format.line_spacing.pt, .79 * 72 + 6, delta=.1)
        self.assertAlmostEqual(d.inline_shapes[0].height / 914400, .79, places=4)
        self.assertAlmostEqual(d.inline_shapes[0].width / d.inline_shapes[0].height, 2, places=4)
        with ZipFile(output) as z:
            embedded = next(n for n in z.namelist() if n.startswith('word/media/'))
            self.assertEqual(z.read(embedded), base64.b64decode(PNG))

    def test_bad_crop_repeat_and_heading_fail_without_output_or_image_changes(self):
        cases = []
        bad = sample_data(); bad['blocks'][1]['images'][0]['crop'] = [0, 0, 999, 120]; cases.append(bad)
        cases.append({'version': 1, 'blocks': [{'repeat_of': 'future'}]})
        cases.append({'version': 1, 'blocks': [{'heading': '标题', 'images': []}]})
        for i, value in enumerate(cases):
            source = self.root / f'bad{i}.json'
            source.write_text(json.dumps(value), encoding='utf-8')
            output = self.root / f'bad{i}.docx'
            with self.assertRaises(ValueError):
                maker.make_handout(source, output, input_format='json')
            self.assertFalse(output.exists())
        self.assertEqual((self.root / 'synthetic.png').read_bytes(), base64.b64decode(PNG))

    def test_json_extension_does_not_enable_new_format(self):
        source = self.root / 'old.json'
        source.write_text('普通旧输入①②++', encoding='utf-8')
        maker.make_handout(source, self.root / 'old.docx')
        self.assertEqual(Document(self.root / 'old.docx').paragraphs[0].text, '普通旧输入①②++')

    def test_audit_distinguishes_error_warning_and_manual_and_is_read_only(self):
        d = Document(); p = d.add_paragraph(); p.add_run('++').italic = True
        p.add_run(' **残留'); f = self.root / 'findings.docx'; d.save(f)
        before = f.read_bytes(); result = inspect_docx(f)
        self.assertTrue(result['errors']); self.assertTrue(result['warnings']); self.assertTrue(result['manual'])
        self.assertEqual(before, f.read_bytes())


class TakeoverTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='sign-update-tests-')
        self.root = Path(self.temp.name) / 'instance'
        (self.root / 'scripts').mkdir(parents=True)
        for name in ('make_handout.py', 'audit_docx.py', 'handout_update.py', 'TOOLING.md'):
            shutil.copyfile(SCRIPTS / name, self.root / 'scripts' / name)
        (self.root / '使用说明.md').write_text('合成使用说明', encoding='utf-8')
        (self.root / 'README.md').write_text('合成下载说明', encoding='utf-8')
        self.original = '# 项目要求\n' + maker.UPDATE_BEGIN + '\n更新说明\n' + maker.UPDATE_END + '\n其余要求必须保留\n'
        (self.root / 'AGENTS.md').write_text(self.original, encoding='utf-8')
        (self.root / '工作台').mkdir()
        self.sentinel = self.root / '工作台' / 'synthetic-sentinel.txt'
        self.sentinel.write_text('合成当前阶段与选择不改动', encoding='utf-8')
        self.state = Path(self.temp.name) / 'local-state' / 'record.json'

    def tearDown(self):
        self.temp.cleanup()

    def takeover(self, action, **kwargs):
        return maker.update_takeover(action, self.root, self.state, **kwargs)

    def test_first_repeat_recover_and_recover_after_recovering_agents(self):
        self.assertEqual(self.takeover('check')['status'], '首次接管')
        self.assertFalse(self.state.exists())
        self.assertEqual((self.root / 'AGENTS.md').read_text(encoding='utf-8'), self.original)
        self.takeover('fail', note='已完成检查；缺依赖待恢复')
        self.assertEqual(self.takeover('check')['status'], '尚未完成')
        self.assertIn(maker.UPDATE_BEGIN, (self.root / 'AGENTS.md').read_text(encoding='utf-8'))
        self.takeover('complete', note='合成验证与页面检查通过；任务与选择未变')
        cleaned = (self.root / 'AGENTS.md').read_text(encoding='utf-8')
        self.assertNotIn(maker.UPDATE_BEGIN, cleaned)
        self.assertIn('其余要求必须保留', cleaned)
        self.assertEqual(self.takeover('check')['status'], '已完成')
        self.assertFalse(self.takeover('complete', note='重复核对成功')['removed_block'])
        (self.root / 'AGENTS.md').write_text(self.original, encoding='utf-8')
        self.assertEqual(self.takeover('check')['status'], '已完成')
        self.assertTrue(self.takeover('complete', note='同版重新覆盖后核对成功')['removed_block'])
        self.assertEqual(self.sentinel.read_text(encoding='utf-8'), '合成当前阶段与选择不改动')
        with (self.root / 'scripts/TOOLING.md').open('a', encoding='utf-8') as f:
            f.write('\n工具说明已变化')
        self.assertEqual(self.takeover('check')['status'], '尚未完成')

    def test_failed_check_never_removes_block_or_records_success(self):
        with patch.object(updater, 'compatibility_check', side_effect=RuntimeError('synthetic failure')):
            with self.assertRaises(RuntimeError):
                self.takeover('complete', note='失败不能清理')
        self.assertFalse(self.state.exists())
        self.assertEqual((self.root / 'AGENTS.md').read_text(encoding='utf-8'), self.original)
        with self.assertRaises(ValueError):
            self.takeover('complete')
        self.assertEqual((self.root / 'AGENTS.md').read_text(encoding='utf-8'), self.original)

    def test_cleanup_failure_preserves_instructions_and_can_resume(self):
        real_replace = Path.replace
        def fail_cleanup(path, target):
            if path.name == '.AGENTS-update.tmp':
                raise OSError('synthetic cleanup failure')
            return real_replace(path, target)
        with patch.object(Path, 'replace', fail_cleanup):
            with self.assertRaises(OSError):
                self.takeover('complete', note='已验证，但清理失败不能宣告完成')
        self.assertEqual((self.root / 'AGENTS.md').read_text(encoding='utf-8'), self.original)
        self.assertEqual(json.loads(self.state.read_text(encoding='utf-8'))['status'], 'incomplete')
        self.assertEqual(self.takeover('check')['status'], '尚未完成')
        self.takeover('complete', note='恢复后再次检查通过')
        self.assertNotIn(maker.UPDATE_BEGIN, (self.root / 'AGENTS.md').read_text(encoding='utf-8'))

    def test_state_write_failure_never_removes_instructions(self):
        with patch.object(updater, 'write_record', side_effect=OSError('synthetic disk failure')):
            with self.assertRaises(OSError):
                self.takeover('complete', note='记录失败')
        self.assertFalse(self.state.exists())
        self.assertEqual((self.root / 'AGENTS.md').read_text(encoding='utf-8'), self.original)

    def test_corrupt_record_and_malformed_block_do_not_hide_failure(self):
        self.state.parent.mkdir(); self.state.write_text('broken', encoding='utf-8')
        self.assertEqual(self.takeover('check')['status'], '首次接管')
        agents = self.root / 'AGENTS.md'
        agents.write_text(maker.UPDATE_BEGIN + '\n缺结束标记', encoding='utf-8')
        with self.assertRaises(ValueError):
            self.takeover('complete', note='不能成功清理残缺区块')
        self.assertIn(maker.UPDATE_BEGIN, agents.read_text(encoding='utf-8'))
        self.assertEqual(self.state.read_text(encoding='utf-8'), 'broken')


if __name__ == '__main__':
    unittest.main()
