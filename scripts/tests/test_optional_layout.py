"""Optional layout checks, synthetic data only."""
import base64
import hashlib
import json
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile
from docx import Document
from test_production import maker, PNG

class OptionalLayoutTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        (self.root/'image.png').write_bytes(base64.b64decode(PNG))
    def tearDown(self):
        self.temp.cleanup()
    def config(self, data):
        path=self.root/'layout.json';path.write_text(json.dumps(data),encoding='utf-8');return path
    def test_roles_and_image_spacing(self):
        source=self.root/'input.json'
        source.write_text(json.dumps({'version':1,'blocks':[
            {'text':'Intro','role':'intro'},
            {'lyrics':'Lyric','transcription':'*action①++*','images':[{'path':'image.png','note':'note'}]}]}),encoding='utf-8')
        config=self.config({'image_space_before':0,'image_space_after':0,'image_line_spacing':1,
            'content_styles':{'intro':{'font_name':'KaiTi','font_size':9},
                              'lyrics':{'font_size':11},'transcription':{'font_size':10.5},
                              'note':{'font_size':9}}})
        out=self.root/'out.docx';maker.make_handout(source,out,layout=config,input_format='json')
        doc=Document(out)
        self.assertEqual(doc.paragraphs[0].runs[0].font.name,'KaiTi')
        self.assertEqual(doc.paragraphs[1].runs[0].font.size.pt,11)
        self.assertEqual(doc.paragraphs[1].runs[1].font.size.pt,10.5)
        self.assertEqual(doc.paragraphs[2].paragraph_format.space_before.pt,0)
        self.assertEqual(doc.paragraphs[2].paragraph_format.space_after.pt,0)
        self.assertTrue(any(r.font.size and r.font.size.pt==9 for r in doc.paragraphs[2].runs))
    def test_padding_only_reference_rows(self):
        source=self.root/'input.json';source.write_text(json.dumps({'version':1,'blocks':[{'text':'x','images':[{'path':'image.png'}]},{'text':'y','images':[{'path':'image.png','reference':True}]}]}),encoding='utf-8')
        out=self.root/'out.docx';maker.make_handout(source,out,layout=self.config({'reference_frame_padding':True,'reference_frame_padding_pt':1}))
        doc=Document(out)
        self.assertEqual(doc.paragraphs[1].paragraph_format.line_spacing,1.25)
        self.assertAlmostEqual(doc.paragraphs[3].paragraph_format.line_spacing.pt,.79*72+1,places=1)
    def test_reference_copy_preserves_parts_and_dimensions(self):
        doc=Document();p=doc.add_paragraph('Original');p.add_run().add_picture(str(self.root/'image.png'))
        source=self.root/'original.docx';doc.save(source)
        before=hashlib.sha256(source.read_bytes()).hexdigest()
        out=self.root/'copy.docx';maker.make_handout(source,out,input_format='docx',layout=self.config({'paragraph_overrides':{'1':{'font_name':'KaiTi','space_after':0}}}))
        with ZipFile(source) as a,ZipFile(out) as b:
            self.assertEqual(a.namelist(),b.namelist())
            for name in a.namelist():
                if name!='word/document.xml':self.assertEqual(a.read(name),b.read(name))
        self.assertEqual(before,hashlib.sha256(source.read_bytes()).hexdigest())
        self.assertEqual(Document(source).inline_shapes[0].height,Document(out).inline_shapes[0].height)
        self.assertEqual(Document(out).paragraphs[0].runs[0].font.name,'KaiTi')
    def test_bad_reference_selector_does_not_create_output(self):
        source=self.root/'original.docx';Document().save(source)
        out=self.root/'copy.docx'
        with self.assertRaises(ValueError):maker.make_handout(source,out,input_format='docx',layout=self.config({'reference_frame_rows':[99]}))
        self.assertFalse(out.exists())
