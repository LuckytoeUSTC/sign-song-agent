"""Current production regression tests; never read real song data."""
import base64
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

SCRIPTS = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SCRIPTS))
import make_handout as maker
from audit_docx import inspect_docx
from docx import Document

PNG = 'iVBORw0KGgoAAAANSUhEUgAAAPAAAAB4CAIAAABD1OhwAAADbklEQVR4nO3d4W6qQBCGYabp/d/yNNHEUFFhhd2Z+fZ9fp5U3aNvp4CI5u4LoOInegHAlQgaUggaUggaUggaUggaUggaUggaUggaUggaUggaUggaUggaUggaUggaUggaUggaUggaUggaUn6bftrMlvr4GOXXr2/+p64taOixliG1/uGccRP0pOz0H9vHPaQq+8ugU/0fptpe6vo8+N7L+vK2939MkgQTeiL2KsemENc//HRvSbLmKMe8NfvN13f48ubhfwaZ0JOmvFzkflfrh4gd1UzouWr2c1P5ne3dRo1qglZm/4+y9Z6aTw8R0jRBy7KgY8axTRO0Jgt9BySwaYIWZAnez4tqmqDVWIKaA5smaCmWpuaopglahyWrOaRpghbkaWoevx6CFpHz3LeHx6p6D2mCVhB+BkWe1RK0FE85nkeujaDLqzWee6+ZoHV44vE8bIUEXVvyfcHxe4cEDSkEXVi58TxgSBM0pBA0pBB0VUW3N3pvdRA0pBA0pBB0SaW3N7pudRA0pBC07PkScyLoozXTdAkEvWN7kStkRtA7klzhSmyPsN9+IUGXbBrvEPQhNF0FQR9F0yUQdAOazm/GC55fuBFsZtX3zMQwoU+h5mwI+hSOeGQz4ybHmbG6/YaHK1aEyzChG1BzfgR9FDWXQNCHUHMVBL2Pmgsh6Ho1D7s0bcWzrAh6x/q5zlAzPiPoffeOqbkEgj6EmqsgaEgh6JIE9gutz+duCBpSCBpSCLqq0lsd1u1zvgQNKQQNKQRdWNGtDut5XRGChhSCrq3ckLbOl30iaB2WvukBKyTo8iqeZ+Ld1kzQUizxkB6zNoJWUGtIe8/VErSI5HuHNuoSwAQtyJI1PXI9BK3DV8MvT9PrlQzYNCJoKZ6s6cE1E7QgT9P0+JoJWlOGpkNqJmhZHtp0VM0ErcyDmg6smaDnato6Z/30ECFv93CUY7rvhbEOWW/vNurNyxkveD4bv7W1/UrcS5rb/nrEvg/PhJ6Fbzo7Oa1f3jz8rBIm9NSjemnch/vwCxCe8l3bt5JlePPpvCRPfSy77qVM9XwyoSflp8/OS9XxA0HPzlsOV+eMeI0vQoUUjnJACkFDCkFDCkFDCkFDCkFDCkFDCkFDCkFDCkFDCkFDCkFDCkFDCkFDCkFDCkFDCkFDCkFjUfIHhfVyD7qKb7kAAAAASUVORK5CYII='
LAYOUT = dict(font_name='SimSun', font_size=12, title_size=18, heading_size=14,
              image_height=.6, margin_horizontal=2.3, margin_vertical=2, line_spacing=1.2)


def sample_data():
    return {'version': 1, 'blocks': [
        {'heading': '合成制作样稿'},
        {'id': 'verse', 'lyrics': '测试画面', 'transcription': '*动作①++*/~~省略②~~',
         'record': '长说明不打印', 'images': [
             {'path': 'synthetic.png', 'entry': '内部词条不打印', 'sense': '义项不打印',
              'source': '来源不打印', 'action': '动作部分不打印', 'record': '图长说明不打印',
              'alt': '替代文本不打印', 'note': '方向调整', 'reference': True,
              'crop': [0, 0, 120, 120]},
             {'path': 'synthetic.png', 'entry': '无标签图', 'note': '双主体位置'}]},
        {'repeat_of': 'verse', 'omit_images': True, 'note': '重复段沿用前图'},
        {'repeat_of': 'verse'}]}


class ProductionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='sign-production-tests-')
        self.root = Path(self.temp.name)
        (self.root / 'synthetic.png').write_bytes(base64.b64decode(PNG))
        (self.root / 'layout.json').write_text(json.dumps(LAYOUT), encoding='utf-8')

    def tearDown(self):
        self.temp.cleanup()

    def run_maker(self, source, output, *args):
        result = subprocess.run([sys.executable, str(SCRIPTS / 'make_handout.py'), str(source),
                                 '--out', str(output), *args], capture_output=True)
        self.assertEqual(result.returncode, 0, result.stderr.decode('utf-8', errors='replace'))

    def test_audit_parses_reused_image_once_per_document(self):
        import audit_docx as checker
        source = self.root / 'repeated.json'
        source.write_text(json.dumps(sample_data()), encoding='utf-8')
        output = self.root / 'repeated.docx'
        maker.make_handout(source, output, input_format='json')
        with patch.object(checker.Image, 'from_file', wraps=checker.Image.from_file) as parse:
            result = checker.inspect_docx(output)
        self.assertEqual(result['images'], 4)
        self.assertEqual(parse.call_count, 1)
        self.assertEqual(result['errors'], [])



    def test_json_notes_crop_repeat_and_markers(self):
        source = self.root / 'content.json'
        source.write_text(json.dumps(sample_data(), ensure_ascii=False), encoding='utf-8')
        before = (self.root / 'synthetic.png').read_bytes()
        output = self.root / 'json.docx'
        self.run_maker(source, output, '--input-format', 'json')
        d = Document(output)
        visible = '\n'.join(p.text for p in d.paragraphs)
        for hidden in ('内部词条不打印', '义项不打印', '来源不打印', '动作部分不打印',
                       '长说明不打印', '替代文本不打印', '无标签图'):
            self.assertNotIn(hidden, visible)
        self.assertEqual(visible.count('测试画面（动作①++/省略②）'), 3)
        self.assertEqual(visible.count('〔方向调整〕'), 2)
        self.assertEqual(visible.count('〔双主体位置〕'), 2)
        self.assertIn('重复段沿用前图', visible)
        self.assertEqual(len(d.inline_shapes), 4)
        self.assertAlmostEqual(d.inline_shapes[0].width / d.inline_shapes[0].height, 1, places=4)
        self.assertAlmostEqual(d.inline_shapes[1].width / d.inline_shapes[1].height, 2, places=4)
        self.assertEqual(before, (self.root / 'synthetic.png').read_bytes())
        with ZipFile(output) as z:
            embedded = next(n for n in z.namelist() if n.startswith('word/media/'))
            self.assertEqual(z.read(embedded), before)
            self.assertIn(b'sysDot', z.read('word/document.xml'))
        flags = [(r.text, r.italic, r.font.strike) for p in d.paragraphs for r in p.runs]
        self.assertIn(('动作①', True, False), flags)
        self.assertIn(('++', False, False), flags)
        self.assertIn(('省略②', False, True), flags)
        result = inspect_docx(output, source=source)
        self.assertFalse(result['errors'], result)
        self.assertFalse(result['warnings'], result)
        d.paragraphs[2].runs[-2].text = '〔误改短注释〕'
        altered = self.root / 'altered.docx'
        d.save(altered)
        self.assertTrue(inspect_docx(altered, source=source)['errors'])

    def test_reference_frame_padding_is_explicit_and_keeps_original_pixels(self):
        from docx.enum.text import WD_LINE_SPACING
        layout = self.root / 'frame.json'
        layout.write_text('{"reference_frame_padding": true}', encoding='utf-8')
        output = self.root / 'frame.docx'
        source = self.root / 'frame.json.input'
        source.write_text(json.dumps({'version': 1, 'blocks': [
            {'text': '参考图', 'images': [{'path': 'synthetic.png', 'reference': True}]}]}), encoding='utf-8')
        maker.make_handout(source, output, layout=layout)
        d = Document(output)
        image_row = next(p for p in d.paragraphs if p._p.xpath('.//w:drawing'))
        self.assertEqual(image_row.paragraph_format.line_spacing_rule, WD_LINE_SPACING.AT_LEAST)
        self.assertAlmostEqual(image_row.paragraph_format.line_spacing.pt, .79 * 72 + 6, delta=.1)
        self.assertAlmostEqual(d.inline_shapes[0].height / 914400, .79, places=4)
        self.assertAlmostEqual(d.inline_shapes[0].width / d.inline_shapes[0].height, 2, places=4)
        with ZipFile(output) as z:
            embedded = next(n for n in z.namelist() if n.startswith('word/media/'))
            self.assertEqual(z.read(embedded), base64.b64decode(PNG))

    def test_bad_crop_repeat_and_heading_fail_without_output_or_image_changes(self):
        cases = []
        bad = sample_data(); bad['blocks'][1]['images'][0]['crop'] = [0, 0, 999, 120]; cases.append(bad)
        cases.append({'version': 1, 'blocks': [{'repeat_of': 'future'}]})
        cases.append({'version': 1, 'blocks': [{'heading': '标题', 'images': []}]})
        for i, value in enumerate(cases):
            source = self.root / f'bad{i}.json'
            source.write_text(json.dumps(value), encoding='utf-8')
            output = self.root / f'bad{i}.docx'
            with self.assertRaises(ValueError):
                maker.make_handout(source, output, input_format='json')
            self.assertFalse(output.exists())
        self.assertEqual((self.root / 'synthetic.png').read_bytes(), base64.b64decode(PNG))


    def test_audit_distinguishes_error_warning_and_manual_and_is_read_only(self):
        d = Document(); p = d.add_paragraph(); p.add_run('++').italic = True
        p.add_run(' **残留'); f = self.root / 'findings.docx'; d.save(f)
        before = f.read_bytes(); result = inspect_docx(f)
        self.assertTrue(result['errors']); self.assertTrue(result['warnings']); self.assertTrue(result['manual'])
        self.assertEqual(before, f.read_bytes())

    def test_default_json_and_removed_legacy_cli(self):
        source = self.root / 'content.json'
        source.write_text(json.dumps(sample_data()), encoding='utf-8')
        self.run_maker(source, self.root / 'default.docx')
        self.assertFalse(inspect_docx(self.root / 'default.docx', source=source)['errors'])
        for args in (['--input-format', 'markdown'], ['--update-action', 'check']):
            result = subprocess.run([sys.executable, str(SCRIPTS / 'make_handout.py'), *args], capture_output=True)
            self.assertNotEqual(result.returncode, 0)
        source.write_text('# old Markdown', encoding='utf-8')
        with self.assertRaises(ValueError):
            maker.make_handout(source, self.root / 'rejected.docx')
        self.assertFalse((self.root / 'rejected.docx').exists())
