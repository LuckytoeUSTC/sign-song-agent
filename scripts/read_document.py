"""Read current source text/formatting and comments; no cache or source edits."""
from pathlib import Path
from zipfile import ZipFile
import argparse
def main():
 ap=argparse.ArgumentParser();ap.add_argument('file',type=Path);a=ap.parse_args();p=a.file
 if p.suffix.lower()=='.md':print(p.read_text(encoding='utf-8'));return
 if p.suffix.lower()=='.pdf':
  import pdfplumber
  with pdfplumber.open(p) as d:
   for i,page in enumerate(d.pages,1):print(f'\n## 第{i}页\n\n{page.extract_text() or "〔本页未提取到文字，请看原页〕"}')
  return
 if p.suffix.lower()!='.docx':raise ValueError('支持DOCX、PDF、Markdown')
 from docx import Document
 from lxml import etree
 d=Document(p)
 for i,paragraph in enumerate(d.paragraphs,1):
  text=''
  for r in paragraph.runs:
   s=r.text
   if r.italic:s='*'+s+'*' if s else s
   if r.font.strike:s='~~'+s+'~~' if s else s
   text+=s
  images=len(paragraph._p.xpath('.//w:drawing'))
  if text or images:print(f'正文第{i}段：{text}'+(f'〔{images}幅图，请看原件〕' if images else ''))
 for i,table in enumerate(d.tables,1):
  print(f'\n表格{i}')
  for row in table.rows:print(' | '.join(c.text for c in row.cells))
 ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
 with ZipFile(p) as z:
  if 'word/comments.xml' in z.namelist():
   doc=etree.fromstring(z.read('word/document.xml'))
   for c in etree.fromstring(z.read('word/comments.xml')):
    cid=c.get('{'+ns['w']+'}id');anchor=doc.xpath('//w:commentRangeStart[@w:id=$id]',namespaces=ns,id=cid)
    context=''.join(anchor[0].getparent().xpath('.//w:t/text()',namespaces=ns)) if anchor else ''
    print(f'\n批注{cid}，{c.get("{"+ns["w"]+"}author", "作者未注明")}：'+''.join(c.xpath('.//w:t/text()',namespaces=ns)))
    print('所在段落：'+context)
if __name__=='__main__':main()
